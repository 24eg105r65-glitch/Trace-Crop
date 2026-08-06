from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

path = r'c:\Users\dines\OneDrive\Desktop\Pushkaran Projects\TraceCrop\TraceCrop_Project_Report_Expanded.docx'
doc = Document(path)
heading = doc.add_heading('Project Screenshots', 1)
heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

img_path = r'c:\Users\dines\OneDrive\Desktop\Pushkaran Projects\TraceCrop\public\report_images\landing_page.png'
if os.path.exists(img_path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('TraceCrop landing page highlighting the main product narrative and solution cards.').italic = True
    doc.add_picture(img_path, width=Inches(6.0))
    doc.add_paragraph()

doc.add_paragraph('Overall, TraceCrop presents a complete digital traceability experience that supports agricultural transparency, multi-role collaboration, and consumer trust.')

doc.save(path)
print('Report updated with screenshots')
